"""CSV and XLSX readers with explicitly selected sheet/header/delimiter."""

import csv
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from academic.errors import ApplicationError, ConfigurationError, ValidationError
from .models import SourceRow, SourceTable

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10000
MAX_COLUMNS = 200


def _check_bytes(content: bytes):
    if not isinstance(content, bytes) or not content:
        raise ValidationError("Wybierz niepusty plik planu.")
    if len(content) > MAX_FILE_BYTES:
        raise ValidationError("Plik planu przekracza 10 MB.")


def _workbook(content: bytes):
    _check_bytes(content)
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 60 * 1024 * 1024:
                raise ValidationError("Rozpakowany arkusz przekracza 60 MB.")
        from openpyxl import load_workbook
        return load_workbook(BytesIO(content), read_only=True, data_only=True)
    except ImportError:
        raise ConfigurationError("Obsługa XLSX wymaga zainstalowania openpyxl z requirements.txt.") from None
    except (BadZipFile, ValueError, KeyError, OSError):
        raise ValidationError("Nie można odczytać pliku XLSX.") from None
    except ApplicationError:
        raise
    except Exception:
        # XML/parser failures raised by openpyxl depend on the selected XML
        # implementation. They are invalid input files, not application crashes.
        raise ValidationError("Nie można odczytać zawartości pliku XLSX.") from None


def sheet_names(content: bytes) -> tuple[str, ...]:
    workbook = _workbook(content)
    try:
        return tuple(workbook.sheetnames)
    finally:
        workbook.close()


def _table(rows, header_row: int) -> SourceTable:
    if not isinstance(header_row, int) or not 1 <= header_row <= MAX_ROWS:
        raise ValidationError("Wybierz poprawny numer wiersza nagłówka.")
    columns = None
    result = []
    for number, row in enumerate(rows, 1):
        if number < header_row:
            continue
        cells = list(row)
        while cells and cells[-1] is None:
            cells.pop()
        if number == header_row:
            if not cells or len(cells) > MAX_COLUMNS:
                raise ValidationError("Nagłówek jest pusty lub zawiera ponad 200 kolumn.")
            columns = tuple(str(value).strip() if value is not None else "" for value in cells)
            if any(not column for column in columns) or len(set(columns)) != len(columns):
                raise ValidationError("Każda kolumna nagłówka musi mieć niepustą, unikalną nazwę.")
            continue
        if not any(value is not None and str(value).strip() for value in cells):
            continue
        if len(cells) > len(columns):
            raise ValidationError(f"Wiersz {number} ma więcej wartości niż kolumny nagłówka.")
        result.append(SourceRow(number, dict(zip(columns, cells + [None] * (len(columns) - len(cells))))))
        if len(result) > MAX_ROWS:
            raise ValidationError("Plan zawiera ponad 10 000 niepustych wierszy.")
    if columns is None:
        raise ValidationError("W pliku nie ma wybranego wiersza nagłówka.")
    if not result:
        raise ValidationError("W pliku nie ma wierszy planu pod wybranym nagłówkiem.")
    return SourceTable(columns, tuple(result))


def read_table(content: bytes, filename: str, *, sheet: str | None = None,
               header_row: int = 1, delimiter: str = ";", encoding: str = "utf-8-sig") -> SourceTable:
    _check_bytes(content)
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        if delimiter not in {";", ",", "\t"} or encoding not in {"utf-8-sig", "cp1250"}:
            raise ValidationError("Wybierz obsługiwany separator i kodowanie CSV.")
        try:
            decoded = content.decode(encoding)
            return _table(csv.reader(StringIO(decoded), delimiter=delimiter, strict=True), header_row)
        except UnicodeDecodeError:
            raise ValidationError("Nie można odczytać CSV w wybranym kodowaniu.") from None
        except csv.Error:
            raise ValidationError("Nieprawidłowy zapis kolumn CSV.") from None
    if suffix == ".xlsx":
        workbook = _workbook(content)
        try:
            if sheet is None or sheet not in workbook.sheetnames:
                raise ValidationError("Wybierz istniejący arkusz XLSX.")
            return _table(workbook[sheet].iter_rows(values_only=True), header_row)
        except ApplicationError:
            raise
        except Exception:
            raise ValidationError("Nie można odczytać wybranego arkusza XLSX.") from None
        finally:
            workbook.close()
    raise ValidationError("Importer obsługuje pliki CSV i XLSX.")
